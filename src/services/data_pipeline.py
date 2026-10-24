from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import shutil
import unicodedata
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from src.services.chunking import TextChunk, _encoding

PIPELINE_VERSION = "docling-v1"
DEFAULT_DATASET_DIR = Path("datasets/docs")
DEFAULT_PROCESSED_DIR = Path("data/processed")
DEFAULT_MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024
DEFAULT_CHUNK_SIZE = 900
DEFAULT_CHUNK_OVERLAP = 120
DEFAULT_MIN_CHUNK_TOKENS = 30
MIN_CLEAN_TEXT_CHARS = 20
SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown"}
# Định dạng văn bản thuần — luôn dùng bộ trích xuất nhẹ (không cần docling), nhanh
# và ổn định cho mọi instance.
PLAIN_TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}


class DatasetPipelineError(ValueError):
    """Base error for dataset processing failures."""


class DatasetValidationError(DatasetPipelineError):
    """Raised when a source file should not enter the RAG pipeline."""


class DuplicateDatasetFileError(DatasetValidationError):
    """Raised when a file duplicates a previous file by content hash."""

    def __init__(self, file_path: Path, duplicate_of: Path) -> None:
        super().__init__(f"{file_path} duplicates {duplicate_of}")
        self.file_path = file_path
        self.duplicate_of = duplicate_of


class DoclingConversionError(DatasetPipelineError):
    """Raised when Docling cannot convert the source document."""


class EmptyDocumentTextError(DatasetPipelineError):
    """Raised when conversion succeeds but produces too little useful text."""


@dataclass(frozen=True)
class ValidatedDocument:
    file_path: Path
    file_hash: str
    file_size_bytes: int
    document_key: str


@dataclass(frozen=True)
class DoclingConversion:
    markdown: str
    json_data: Any
    metadata: dict[str, Any]
    image_paths: list[str]


@dataclass(frozen=True)
class ConvertedDocument:
    raw_markdown_path: Path
    raw_json_path: Path
    docling_metadata_path: Path
    markdown: str
    json_data: Any
    metadata: dict[str, Any]
    image_paths: list[str]


@dataclass(frozen=True)
class DocumentSection:
    title: str
    heading_level: int
    content: str
    page_start: int | None
    page_end: int | None
    block_types: list[str]


@dataclass(frozen=True)
class ProcessedDocument:
    validation: ValidatedDocument
    processed_dir: Path
    raw_markdown_path: Path
    raw_json_path: Path
    clean_markdown_path: Path
    metadata_path: Path
    report_path: Path
    document_metadata: dict[str, Any]
    chunks: list[TextChunk]


DoclingRunner = Callable[[Path, Path], DoclingConversion]


def collect_dataset_documents(dataset_dir: Path = DEFAULT_DATASET_DIR) -> list[Path]:
    dataset_path = Path(dataset_dir)
    if not dataset_path.exists():
        return []
    return sorted(
        path
        for path in dataset_path.iterdir()
        if path.is_file() and path.suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS
    )


def collect_dataset_pdfs(dataset_dir: Path = DEFAULT_DATASET_DIR) -> list[Path]:
    return collect_dataset_documents(dataset_dir)


def validate_document_file(
    file_path: Path,
    *,
    seen_hashes: dict[str, Path] | None = None,
    max_file_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
) -> ValidatedDocument:
    path = Path(file_path)
    if not path.exists():
        raise DatasetValidationError(f"File does not exist: {path}")
    if not path.is_file():
        raise DatasetValidationError(f"Path is not a file: {path}")
    if path.suffix.lower() not in SUPPORTED_DOCUMENT_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
        raise DatasetValidationError(f"Only document files are supported ({supported}): {path}")

    file_size_bytes = path.stat().st_size
    if file_size_bytes > max_file_size_bytes:
        raise DatasetValidationError(
            f"File exceeds {max_file_size_bytes} bytes: {path} ({file_size_bytes} bytes)"
        )

    file_hash = calculate_sha256(path)
    if seen_hashes is not None:
        duplicate_of = seen_hashes.get(file_hash)
        if duplicate_of is not None and duplicate_of != path:
            raise DuplicateDatasetFileError(path, duplicate_of)
        seen_hashes[file_hash] = path

    return ValidatedDocument(
        file_path=path,
        file_hash=file_hash,
        file_size_bytes=file_size_bytes,
        document_key=_document_key(path, file_hash),
    )


def validate_pdf_file(
    file_path: Path,
    *,
    seen_hashes: dict[str, Path] | None = None,
    max_file_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
) -> ValidatedDocument:
    return validate_document_file(
        file_path,
        seen_hashes=seen_hashes,
        max_file_size_bytes=max_file_size_bytes,
    )


def calculate_sha256(file_path: Path) -> str:
    digest = hashlib.sha256()
    with Path(file_path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def convert_with_docling(
    file_path: Path,
    processed_dir: Path,
    *,
    docling_runner: DoclingRunner | None = None,
    use_cache: bool = True,
) -> ConvertedDocument:
    processed_dir.mkdir(parents=True, exist_ok=True)
    raw_markdown_path = processed_dir / "raw.md"
    raw_json_path = processed_dir / "docling.json"
    docling_metadata_path = processed_dir / "docling_metadata.json"

    if use_cache and raw_markdown_path.exists() and raw_json_path.exists() and docling_metadata_path.exists():
        _normalize_image_artifacts(processed_dir)
        metadata = dict(_read_json(docling_metadata_path) or {})
        image_paths = sorted(str(path) for path in metadata.get("image_paths", []))
        metadata["image_paths"] = image_paths
        return ConvertedDocument(
            raw_markdown_path=raw_markdown_path,
            raw_json_path=raw_json_path,
            docling_metadata_path=docling_metadata_path,
            markdown=raw_markdown_path.read_text(encoding="utf-8"),
            json_data=_read_json(raw_json_path),
            metadata=metadata,
            image_paths=image_paths,
        )

    try:
        conversion = docling_runner(file_path, processed_dir) if docling_runner else _run_docling_document(file_path, processed_dir)
    except Exception as exc:
        raise DoclingConversionError(str(exc)) from exc

    raw_markdown_path.write_text(conversion.markdown, encoding="utf-8")
    _write_json(raw_json_path, conversion.json_data)
    _write_json(docling_metadata_path, conversion.metadata)
    _normalize_image_artifacts(processed_dir)
    metadata = dict(_read_json(docling_metadata_path) or {})
    image_paths = sorted(str(path) for path in metadata.get("image_paths", conversion.image_paths))
    metadata["image_paths"] = image_paths
    return ConvertedDocument(
        raw_markdown_path=raw_markdown_path,
        raw_json_path=raw_json_path,
        docling_metadata_path=docling_metadata_path,
        markdown=raw_markdown_path.read_text(encoding="utf-8"),
        json_data=_read_json(raw_json_path),
        metadata=metadata,
        image_paths=image_paths,
    )


def _extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    # Nối bằng \f (form feed) để lightweight_convert đếm ĐÚNG số trang
    # (page_count = markdown.count("\f") + 1). Trước đây nối "\n\n" → mọi PDF bị
    # ghi total_pages=1 → chatbot trả lời sai số trang tài liệu.
    return "\f".join(pages)


def _extract_docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        lines.append(f"## {text}" if style.startswith("heading") else text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n\n".join(lines)


def _extract_pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        texts.append(text)
        if texts:
            parts.append(f"## Slide {index}\n\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_txt_text(path: Path) -> str:
    """Đọc file văn bản thuần (.txt/.md) — thử vài encoding phổ biến cho ổn định."""
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding).strip()
        except (UnicodeDecodeError, LookupError):
            continue
    # Phương án cuối: bỏ qua ký tự lỗi để không chết toàn bộ tài liệu.
    return raw.decode("utf-8", errors="ignore").strip()


def extract_text_quick(file_path: Path, *, max_chars: int = 6000, max_pdf_pages: int = 6) -> str:
    """Trích xuất NHANH một phần văn bản từ file để AI gợi ý tiêu đề/mô tả.

    Giới hạn số trang PDF + số ký tự để phản hồi nhanh (không đọc cả tài liệu lớn).
    Best-effort: lỗi định dạng → trả "" để không chặn luồng.
    """
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            parts = []
            for page in reader.pages[:max_pdf_pages]:
                text = (page.extract_text() or "").strip()
                if text:
                    parts.append(text)
            text = "\n\n".join(parts)
        elif suffix == ".docx":
            text = _extract_docx_text(file_path)
        elif suffix == ".pptx":
            text = _extract_pptx_text(file_path)
        elif suffix in PLAIN_TEXT_EXTENSIONS:
            text = _extract_txt_text(file_path)
        else:
            return ""
    except Exception:
        return ""
    return (text or "").strip()[:max_chars]


def lightweight_convert(file_path: Path, processed_dir: Path) -> DoclingConversion:
    """Trích xuất văn bản nhẹ (pypdf/python-docx/python-pptx/text) — KHÔNG cần docling/torch.

    Dùng khi USE_DOCLING=false (instance RAM thấp như Render free) để ingest không bị
    OOM, và luôn dùng cho file văn bản thuần (.txt/.md). Bỏ qua bảng/hình phức tạp;
    đủ dùng cho RAG hỏi-đáp văn bản.
    """
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        markdown = _extract_pdf_text(file_path)
    elif suffix == ".docx":
        markdown = _extract_docx_text(file_path)
    elif suffix == ".pptx":
        markdown = _extract_pptx_text(file_path)
    elif suffix in PLAIN_TEXT_EXTENSIONS:
        markdown = _extract_txt_text(file_path)
    else:
        raise DoclingConversionError(f"Lightweight extractor không hỗ trợ định dạng: {suffix}")

    page_count = markdown.count("\f") + 1 if markdown else 0
    return DoclingConversion(
        markdown=markdown,
        json_data=None,
        metadata={"page_count": page_count, "parser": "lightweight"},
        image_paths=[],
    )


def clean_markdown(markdown: str) -> str:
    # 1. Unicode NFC normalization
    text = unicodedata.normalize("NFC", markdown)

    # 2. HTML Entity Unescaping
    text = html.unescape(text)

    # 3. Carriage return normalization
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # \f (form feed) ch\u1ec9 d\u00f9ng \u0111\u1ec3 \u0111\u1ebfm s\u1ed1 trang \u1edf b\u01b0\u1edbc convert \u2014 lo\u1ea1i kh\u1ecfi n\u1ed9i dung chunk.
    text = text.replace("\f", "\n\n")
    text = text.replace("\u00a0", " ")

    # 4. Clean backslash-escaped underscores
    text = text.replace(r"\_", "_")

    # 4b. Drop image references entirely: image files are tracked in metadata,
    # keeping the links inline only pollutes embeddings and section titles.
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)

    # 5. Spacing normalization around punctuation
    # Remove space before punctuation (horizontal whitespace only — matching "\s"
    # here would glue the next line onto the previous sentence)
    text = re.sub(r"[ \t]+([,.!?;:])", r"\1", text)
    text = _add_space_after_punctuation(text)

    # 6. Remove empty/malformed code blocks (e.g. ```\nShell\n``` or empty ```\n```)
    text = re.sub(r"(?m)^```\w*\s*\n\s*```$", "", text)
    lang_pattern = r"(Shell|Python|Public|None|Html|Css|Javascript|Js|Bash|C\+\+|Java|C#|Ruby|Markdown|Yaml|Json|Sql)"
    text = re.sub(
        r"(?m)^```\w*\n\s*" + lang_pattern + r"\s*\n\s*```$",
        "",
        text,
        flags=re.IGNORECASE
    )

    # 7. Remove line hyphenation
    text = re.sub(r"(?<=\w)-\n(?=\w)", "", text)

    # Collapse multiple spaces/tabs
    text = re.sub(r"[ \t]+", " ", text)

    lines = [line.rstrip() for line in text.split("\n")]
    noise_counts = Counter(
        candidate for line in lines if (candidate := _noise_candidate(line))
    )
    filtered_lines = [
        line
        for line in lines
        if not _is_page_marker(line)
        and not (
            (candidate := _noise_candidate(line))
            and noise_counts[candidate] >= 3
        )
    ]
    cleaned = _join_markdown_lines(filtered_lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


_URL_PATTERN = re.compile(r"(?:https?://|www\.)\S+")


def _add_space_after_punctuation(text: str) -> str:
    """Add a space after sentence punctuation, leaving URLs untouched."""

    def fix(segment: str) -> str:
        # Skip digits (3.7), paths, and markdown syntax like "![" / "["
        segment = re.sub(r"(?<!\d)([,;!?:])(?!\d|\s|[\[\(/\\])", r"\1 ", segment)
        return re.sub(r"(?<!\d)\.(?=[A-ZÀ-Ỹ])", r". ", segment)

    parts: list[str] = []
    last = 0
    for match in _URL_PATTERN.finditer(text):
        parts.append(fix(text[last : match.start()]))
        parts.append(match.group(0))
        last = match.end()
    parts.append(fix(text[last:]))
    return "".join(parts)


def _extract_sections(json_data: Any, docling_metadata: dict[str, Any], markdown: str) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []

    # 1. Try to extract from json_data if it is a dict with "texts"
    if isinstance(json_data, dict) and "texts" in json_data:
        texts = json_data.get("texts", [])
        for text_item in texts:
            if text_item.get("label") == "section_header":
                title = text_item.get("text", "").strip()
                level = text_item.get("level", 1)
                prov = text_item.get("prov", [])
                page_start = None
                if prov and isinstance(prov, list):
                    page_start = prov[0].get("page_no")
                if title:
                    sections.append({
                        "title": title,
                        "heading_level": int(level),
                        "page_start": int(page_start) if page_start is not None else None,
                    })

    # 2. Try to fallback to table_of_contents in docling_metadata
    if not sections and isinstance(docling_metadata, dict):
        toc = docling_metadata.get("table_of_contents") or []
        for item in toc:
            title = str(item.get("title") or "").strip()
            if not title:
                continue
            page_id = item.get("page_id")
            sections.append({
                "title": title,
                "heading_level": int(item.get("heading_level") or 1),
                "page_start": int(page_id) + 1 if isinstance(page_id, int) else None,
            })

    # 3. Try fallback to parsing markdown
    if not sections:
        for title, heading_level, _ in _iter_heading_sections(markdown):
            sections.append({
                "title": title,
                "heading_level": heading_level,
                "page_start": None,
            })

    return sections


def _extract_tables_and_images(json_data: Any, markdown: str) -> tuple[int, int, list[dict[str, Any]], list[dict[str, Any]]]:
    tables_count = 0
    images_count = 0
    tables: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []

    # Extract from json_data if it is a dict
    if isinstance(json_data, dict):
        # Tables
        raw_tables = json_data.get("tables", [])
        tables_count = len(raw_tables)
        for idx, tbl in enumerate(raw_tables, start=1):
            prov = tbl.get("prov", [])
            page_no = prov[0].get("page_no") if prov and isinstance(prov, list) else None
            tables.append({
                "index": idx,
                "page_no": page_no,
                "caption": tbl.get("captions", [{}])[0].get("text") if tbl.get("captions") else None
            })

        # Pictures / Images
        raw_pictures = json_data.get("pictures", [])
        images_count = len(raw_pictures)
        for idx, pic in enumerate(raw_pictures, start=1):
            prov = pic.get("prov", [])
            page_no = prov[0].get("page_no") if prov and isinstance(prov, list) else None
            img_info = pic.get("image", {})
            uri = img_info.get("uri")
            images.append({
                "index": idx,
                "page_no": page_no,
                "uri": uri
            })
    else:
        # Fallback using regex / string search in markdown if json_data is not available/parsed
        table_separator_count = len(re.findall(r"\n\|[- :]+\|", markdown))
        tables_count = table_separator_count

        markdown_images = re.findall(r"!\[(.*?)\]\((.*?)\)", markdown)
        images_count = len(markdown_images)
        for idx, (caption, uri) in enumerate(markdown_images, start=1):
            images.append({
                "index": idx,
                "page_no": None,
                "uri": uri
            })

    return tables_count, images_count, tables, images


def build_document_metadata(
    *,
    validation: ValidatedDocument,
    converted: ConvertedDocument,
    clean_markdown_text: str,
    processed_dir: Path,
    dataset_name: str,
) -> dict[str, Any]:
    block_counts = _count_docling_blocks(converted.json_data)
    page_count = _page_count(converted.json_data, converted.metadata)

    sections = _extract_sections(converted.json_data, converted.metadata, clean_markdown_text)
    tables_count, images_count, tables, images = _extract_tables_and_images(converted.json_data, clean_markdown_text)

    language = detect_language(clean_markdown_text)

    return {
        "dataset": dataset_name,
        "source_path": str(validation.file_path),
        "file_type": validation.file_path.suffix.lstrip(".").lower(),
        "file_hash": validation.file_hash,
        "file_size_bytes": validation.file_size_bytes,
        "processed_dir": str(processed_dir),
        "parser": "docling",
        "pipeline_version": PIPELINE_VERSION,
        "language": language,
        "page_count": page_count,
        "section_count": len(sections),
        "sections": sections,
        "tables_count": tables_count,
        "images_count": images_count,
        "tables": tables,
        "images": images,
        "processing_status": "ready",
        "processing_error": "",
        "extraction_metadata": {
            "raw_markdown_path": str(converted.raw_markdown_path),
            "raw_json_path": str(converted.raw_json_path),
            "docling_metadata_path": str(converted.docling_metadata_path),
            "image_paths": converted.image_paths,
            "block_counts": block_counts,
            "page_count": page_count,
            "conversion_status": converted.metadata.get("conversion_status", ""),
            "errors": converted.metadata.get("errors", []),
            "timings": converted.metadata.get("timings", {}),
        },
        "customer_metadata": {
            "source": "batch_local",
            "uploaded_by": None,
        },
    }


def chunk_clean_markdown(
    clean_markdown_text: str,
    *,
    marker_json: Any,
    document_metadata: dict[str, Any],
    model_name: str,
    source_markdown_path: Path,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
    min_chunk_tokens: int = DEFAULT_MIN_CHUNK_TOKENS,
) -> list[TextChunk]:
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    encoding = _encoding(model_name)
    sections = sections_from_markdown(
        clean_markdown_text,
        marker_json=marker_json,
        document_metadata=document_metadata,
    )
    chunks: list[TextChunk] = []

    for section in sections:
        units = _split_markdown_units(section.content)
        current_units: list[str] = []
        current_tokens = 0
        overlap_text = ""

        for unit in units:
            unit_tokens = len(encoding.encode(unit))
            if unit_tokens > chunk_size:
                if current_units:
                    overlap_text = _append_structured_chunk(
                        chunks,
                        current_units,
                        section,
                        encoding,
                        chunk_overlap,
                        source_markdown_path,
                    )
                    current_units = []
                    current_tokens = 0
                for token_slice in _sliding_token_slices(
                    encoding.encode(unit),
                    chunk_size=chunk_size,
                    chunk_overlap=chunk_overlap,
                ):
                    content = encoding.decode(token_slice).strip()
                    if content:
                        _append_text_chunk(
                            chunks,
                            content,
                            section,
                            len(token_slice),
                            source_markdown_path,
                        )
                overlap_text = ""
                continue

            candidate_tokens = current_tokens + unit_tokens
            if current_units and candidate_tokens > chunk_size:
                overlap_text = _append_structured_chunk(
                    chunks,
                    current_units,
                    section,
                    encoding,
                    chunk_overlap,
                    source_markdown_path,
                )
                current_units = []
                current_tokens = 0
                if overlap_text:
                    overlap_tokens = len(encoding.encode(overlap_text))
                    if overlap_tokens + unit_tokens <= chunk_size:
                        current_units.append(overlap_text)
                        current_tokens += overlap_tokens

            current_units.append(unit)
            current_tokens += unit_tokens

        if current_units:
            _append_structured_chunk(
                chunks,
                current_units,
                section,
                encoding,
                chunk_overlap,
                source_markdown_path,
            )

    chunks = _merge_small_chunks(
        chunks,
        encoding,
        min_tokens=min(min_chunk_tokens, chunk_size // 4),
        chunk_size=chunk_size,
    )

    return [
        TextChunk(
            chunk_index=index,
            content=chunk.content,
            page_number=chunk.page_number,
            token_count=chunk.token_count,
            metadata={**chunk.metadata, "chunk_index": index},
        )
        for index, chunk in enumerate(chunks)
    ]


def _assign_pages_from_pagebreaks(chunks: list[TextChunk], raw_markdown: str) -> list[TextChunk]:
    """Gán SỐ TRANG thật cho chunk dựa trên ranh giới trang \\f trong văn bản thô.

    Bộ trích xuất nhẹ (pypdf) nối các trang bằng \\f; clean_markdown loại \\f trước khi
    chunk nên chunk mất thông tin trang. Hàm này khớp lại vị trí từng chunk trong văn
    bản thô (so khớp đã chuẩn hoá, bỏ khoảng trắng/ký tự đặc biệt) → trang = số \\f đứng
    trước + 1. Best-effort: chunk không khớp được thì giữ nguyên trang cũ.
    """
    if "\f" not in (raw_markdown or "") or not chunks:
        return chunks
    import bisect
    from dataclasses import replace

    def norm(s: str) -> str:
        return re.sub(r"[\W_]+", "", (s or "").lower(), flags=re.UNICODE)

    page_texts = raw_markdown.split("\f")
    norm_pages = [norm(p) for p in page_texts]
    full = "".join(norm_pages)
    ends, acc = [], 0
    for np in norm_pages:
        acc += len(np)
        ends.append(acc)

    out: list[TextChunk] = []
    search_from = 0
    for chunk in chunks:
        probe = norm(chunk.content)[:60]
        page = None
        if len(probe) >= 20:
            idx = full.find(probe, search_from)
            if idx < 0:
                idx = full.find(probe)
            if idx >= 0:
                page = bisect.bisect_right(ends, idx) + 1
                page = min(page, len(page_texts))
                search_from = idx
        if page:
            out.append(
                replace(chunk, page_number=page, metadata={**chunk.metadata, "page_start": page})
            )
        else:
            out.append(chunk)
    return out


def process_document_for_rag(
    file_path: Path,
    *,
    dataset_name: str = "docs",
    processed_root: Path = DEFAULT_PROCESSED_DIR,
    model_name: str,
    seen_hashes: dict[str, Path] | None = None,
    validation: ValidatedDocument | None = None,
    docling_runner: DoclingRunner | None = None,
    use_cache: bool = True,
) -> ProcessedDocument:
    validation = validation or validate_document_file(file_path, seen_hashes=seen_hashes)
    processed_dir = processed_root / validation.document_key
    converted = convert_with_docling(
        validation.file_path,
        processed_dir,
        docling_runner=docling_runner,
        use_cache=use_cache,
    )
    clean_text = clean_markdown(converted.markdown)
    if len(_visible_text(clean_text)) < MIN_CLEAN_TEXT_CHARS:
        _write_json(
            processed_dir / "pipeline_report.json",
            {
                "status": "failed_empty_text",
                "error": "Converted markdown has too little usable text.",
                "source_path": str(validation.file_path),
                "file_hash": validation.file_hash,
            },
        )
        raise EmptyDocumentTextError("Converted markdown has too little usable text.")

    clean_markdown_path = processed_dir / "clean.md"
    metadata_path = processed_dir / "metadata.json"
    report_path = processed_dir / "pipeline_report.json"
    clean_markdown_path.write_text(clean_text, encoding="utf-8")

    document_metadata = build_document_metadata(
        validation=validation,
        converted=converted,
        clean_markdown_text=clean_text,
        processed_dir=processed_dir,
        dataset_name=dataset_name,
    )
    chunks = chunk_clean_markdown(
        clean_text,
        marker_json=converted.json_data,
        document_metadata=document_metadata,
        model_name=model_name,
        source_markdown_path=clean_markdown_path,
    )
    # PDF trích xuất nhẹ: khớp lại từng chunk với ranh giới trang \f để citation
    # trỏ ĐÚNG trang (tr.N) thay vì mặc định trang 1/None.
    chunks = _assign_pages_from_pagebreaks(chunks, converted.markdown)
    document_metadata["chunk_count"] = len(chunks)

    _write_json(metadata_path, document_metadata)
    _write_json(
        report_path,
        {
            "status": "ready",
            "source_path": str(validation.file_path),
            "file_hash": validation.file_hash,
            "chunk_count": len(chunks),
            "processed_dir": str(processed_dir),
        },
    )
    return ProcessedDocument(
        validation=validation,
        processed_dir=processed_dir,
        raw_markdown_path=converted.raw_markdown_path,
        raw_json_path=converted.raw_json_path,
        clean_markdown_path=clean_markdown_path,
        metadata_path=metadata_path,
        report_path=report_path,
        document_metadata=document_metadata,
        chunks=chunks,
    )


def process_pdf_for_rag(
    file_path: Path,
    *,
    dataset_name: str = "docs",
    processed_root: Path = DEFAULT_PROCESSED_DIR,
    model_name: str,
    seen_hashes: dict[str, Path] | None = None,
    validation: ValidatedDocument | None = None,
    marker_runner: DoclingRunner | None = None,
    docling_runner: DoclingRunner | None = None,
    use_cache: bool = True,
) -> ProcessedDocument:
    return process_document_for_rag(
        file_path,
        dataset_name=dataset_name,
        processed_root=processed_root,
        model_name=model_name,
        seen_hashes=seen_hashes,
        validation=validation,
        docling_runner=docling_runner or marker_runner,
        use_cache=use_cache,
    )


def sections_from_markdown(
    markdown: str,
    *,
    marker_json: Any,
    document_metadata: dict[str, Any],
) -> list[DocumentSection]:
    heading_matches = list(_iter_heading_sections(markdown))
    block_types = sorted(_count_docling_blocks(marker_json))
    page_count = document_metadata.get("page_count") or None
    section_pages = {
        _normalize_title(item.get("title", "")): int(item["page_start"])
        for item in document_metadata.get("sections", [])
        if item.get("title") and item.get("page_start")
    }

    if not heading_matches:
        return [
            DocumentSection(
                title="Document",
                heading_level=0,
                content=markdown,
                page_start=1 if page_count else None,
                page_end=int(page_count) if page_count else None,
                block_types=block_types,
            )
        ]

    sections: list[DocumentSection] = []
    for index, (title, heading_level, content) in enumerate(heading_matches):
        next_title = heading_matches[index + 1][0] if index + 1 < len(heading_matches) else None
        page_start = section_pages.get(_normalize_title(title))
        next_page_start = section_pages.get(_normalize_title(next_title or ""))
        page_end = next_page_start if next_page_start else page_count
        sections.append(
            DocumentSection(
                title=title,
                heading_level=heading_level,
                content=content,
                page_start=page_start,
                page_end=int(page_end) if page_end else page_start,
                block_types=block_types,
            )
        )
    return sections


def detect_language(text: str) -> str:
    sample = text[:5000].lower()

    # Count Vietnamese diacritics
    vi_diacritics = r"[ăâđêôơưáàảãạấầẩẫậắằẳẵặéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵđ]"
    vi_diacritics_count = len(re.findall(vi_diacritics, sample))

    # Vietnamese stopwords/common words
    vi_stopwords = {
        "và", "của", "là", "trong", "cho", "để", "có", "các", "những",
        "người", "được", "một", "này", "khi", "đến", "với", "theo",
        "tại", "nếu", "như", "nhưng", "đã", "sẽ", "đang", "đó", "về"
    }
    # English stopwords/common words
    en_stopwords = {
        "the", "of", "and", "to", "a", "in", "is", "that", "for",
        "it", "on", "with", "as", "this", "are", "by", "an", "be",
        "from", "at", "or", "which", "have", "not", "but"
    }

    # Tokenize words (alphanumeric only)
    words = re.findall(r"\b[a-zà-ỹđ]{2,}\b", sample)

    vi_words_count = sum(1 for w in words if w in vi_stopwords)
    en_words_count = sum(1 for w in words if w in en_stopwords)

    # Scoring
    vi_score = vi_diacritics_count * 2.0 + vi_words_count * 5.0
    en_score = en_words_count * 5.0

    if vi_score > 10.0 and vi_score > en_score:
        return "vi"
    if en_score > 10.0:
        return "en"

    # Fallback to simple diacritic check
    if vi_diacritics_count > 0:
        return "vi"
    if re.search(r"[a-z]", sample):
        return "en"

    return "unknown"


def _run_docling_document(file_path: Path, processed_dir: Path) -> DoclingConversion:
    try:
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions
        from docling.document_converter import PdfFormatOption
        from docling_core.types.doc import ImageRefMode
    except ImportError as exc:
        raise DoclingConversionError(
            "docling is not installed. Install dependencies from requirements.txt in the conda venv first."
        ) from exc

    images_dir = processed_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    converter = _build_docling_converter(file_path, InputFormat, PdfPipelineOptions, PdfFormatOption)
    conversion_result = converter.convert(file_path)
    document = conversion_result.document
    raw_markdown_path = processed_dir / "raw.md"
    raw_json_path = processed_dir / "docling.json"

    # artifacts_dir must be absolute: docling resolves relative paths against the
    # output file's parent, which used to create a stray nested copy of every image.
    document.save_as_markdown(
        raw_markdown_path,
        artifacts_dir=images_dir.resolve(),
        image_mode=ImageRefMode.REFERENCED,
    )
    document.save_as_json(
        raw_json_path,
        artifacts_dir=images_dir.resolve(),
        image_mode=ImageRefMode.REFERENCED,
    )

    image_paths = _find_image_paths(images_dir, processed_dir)
    metadata = _docling_conversion_metadata(conversion_result, image_paths)
    return DoclingConversion(
        markdown=raw_markdown_path.read_text(encoding="utf-8"),
        json_data=_read_json(raw_json_path),
        metadata=metadata,
        image_paths=image_paths,
    )


def _build_docling_converter(
    file_path: Path,
    input_format: Any,
    pdf_pipeline_options: Any,
    pdf_format_option: Any,
) -> Any:
    from docling.document_converter import DocumentConverter

    if file_path.suffix.lower() != ".pdf":
        return DocumentConverter()

    pipeline_options = pdf_pipeline_options()
    pipeline_options.do_table_structure = True
    pipeline_options.generate_page_images = True
    pipeline_options.generate_picture_images = True
    pipeline_options.images_scale = 2.0
    return DocumentConverter(
        format_options={
            input_format.PDF: pdf_format_option(pipeline_options=pipeline_options)
        }
    )


def _normalize_image_artifacts(processed_dir: Path) -> None:
    """Keep exactly one copy of every image under <processed_dir>/images.

    Older pipeline runs left two problems behind: docling wrote its referenced
    artifacts into a stray nested ``<processed_dir>/data/processed/...`` tree,
    and a second export duplicated every page render as ``page-NNN.png``. This
    migrates stray files into ``images/``, drops the duplicate renders, and
    rewrites references in raw.md / docling.json to plain ``images/...`` paths.
    """
    images_dir = processed_dir / "images"
    stray_root = processed_dir / "data"
    if stray_root.exists():
        images_dir.mkdir(parents=True, exist_ok=True)
        for path in sorted(stray_root.rglob("*")):
            if not path.is_file():
                continue
            target = images_dir / path.name
            if target.exists():
                path.unlink()
            else:
                path.replace(target)
        shutil.rmtree(stray_root, ignore_errors=True)

    if images_dir.exists():
        for path in images_dir.iterdir():
            if re.fullmatch(r"(page|picture|table)-\d+\.(png|jpg|jpeg|webp)", path.name):
                path.unlink()

    image_prefixes = {
        f"{processed_dir.as_posix()}/images/",
        f"{(processed_dir / 'images').resolve().as_posix()}/",
        f"data/processed/{processed_dir.name}/images/",
    }
    for artifact_name in ("raw.md", "docling.json"):
        artifact = processed_dir / artifact_name
        if not artifact.exists():
            continue
        text = artifact.read_text(encoding="utf-8")
        updated = text
        for prefix in image_prefixes:
            updated = updated.replace(prefix, "images/")
        if updated != text:
            artifact.write_text(updated, encoding="utf-8")

    metadata_path = processed_dir / "docling_metadata.json"
    if metadata_path.exists() and images_dir.exists():
        metadata = _read_json(metadata_path)
        if isinstance(metadata, dict):
            image_paths = _find_image_paths(images_dir, processed_dir)
            if metadata.get("image_paths") != image_paths:
                metadata["image_paths"] = image_paths
                _write_json(metadata_path, metadata)


def _find_image_paths(images_dir: Path, processed_dir: Path) -> list[str]:
    if not images_dir.exists():
        return []
    return [
        str(path.relative_to(processed_dir))
        for path in sorted(images_dir.rglob("*"))
        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
    ]


def _docling_conversion_metadata(conversion_result: Any, image_paths: list[str]) -> dict[str, Any]:
    document = conversion_result.document
    pages = getattr(document, "pages", {}) or {}
    status = getattr(conversion_result, "status", "")
    errors = getattr(conversion_result, "errors", []) or []
    timings = getattr(conversion_result, "timings", {}) or {}
    return {
        "conversion_status": str(status),
        "errors": _json_safe(errors),
        "timings": _json_safe(timings),
        "page_count": len(pages),
        "image_paths": image_paths,
    }


def _merge_small_chunks(
    chunks: list[TextChunk],
    encoding: Any,
    *,
    min_tokens: int,
    chunk_size: int,
) -> list[TextChunk]:
    if min_tokens <= 0 or len(chunks) < 2:
        return chunks

    def combine(first: TextChunk, second: TextChunk) -> TextChunk:
        content = f"{first.content}\n\n{second.content}".strip()
        # The larger chunk carries the meaningful section metadata; the small
        # one is usually a bare slide title.
        base = first if first.token_count >= second.token_count else second
        return TextChunk(
            chunk_index=second.chunk_index,
            content=content,
            page_number=first.page_number or second.page_number,
            token_count=len(encoding.encode(content)),
            metadata={
                **base.metadata,
                "page_start": first.metadata.get("page_start") or second.metadata.get("page_start"),
                "content_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            },
        )

    merged: list[TextChunk] = []
    pending: TextChunk | None = None
    for chunk in chunks:
        if pending is not None:
            candidate = combine(pending, chunk)
            if candidate.token_count <= chunk_size:
                chunk = candidate
            else:
                merged.append(pending)
            pending = None
        if chunk.token_count < min_tokens:
            pending = chunk
            continue
        merged.append(chunk)

    if pending is not None:
        if merged:
            candidate = combine(merged[-1], pending)
            if candidate.token_count <= chunk_size:
                merged[-1] = candidate
            else:
                merged.append(pending)
        else:
            merged.append(pending)
    return merged


def _append_structured_chunk(
    chunks: list[TextChunk],
    units: list[str],
    section: DocumentSection,
    encoding: Any,
    chunk_overlap: int,
    source_markdown_path: Path,
) -> str:
    content = "\n\n".join(unit.strip() for unit in units if unit.strip()).strip()
    if not content:
        return ""
    tokens = encoding.encode(content)
    _append_text_chunk(chunks, content, section, len(tokens), source_markdown_path)
    if chunk_overlap <= 0:
        return ""
    return encoding.decode(tokens[-chunk_overlap:]).strip()


def _append_text_chunk(
    chunks: list[TextChunk],
    content: str,
    section: DocumentSection,
    token_count: int,
    source_markdown_path: Path,
) -> None:
    metadata = {
        "page_start": section.page_start,
        "page_end": section.page_end,
        "section_title": section.title,
        "heading_level": section.heading_level,
        "block_types": section.block_types,
        "source_markdown_path": str(source_markdown_path),
        "content_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
    }
    chunks.append(
        TextChunk(
            chunk_index=len(chunks),
            content=content,
            page_number=section.page_start,
            token_count=token_count,
            metadata=metadata,
        )
    )


def _sliding_token_slices(
    tokens: list[Any],
    *,
    chunk_size: int,
    chunk_overlap: int,
) -> list[list[Any]]:
    step = chunk_size - chunk_overlap
    slices: list[list[Any]] = []
    for start in range(0, len(tokens), step):
        slices.append(tokens[start : start + chunk_size])
        if start + chunk_size >= len(tokens):
            break
    return slices


def _split_markdown_units(markdown: str) -> list[str]:
    units: list[str] = []
    buffer: list[str] = []
    in_code = False
    in_table = False

    def flush() -> None:
        nonlocal buffer
        content = "\n".join(buffer).strip()
        if content:
            units.append(content)
        buffer = []

    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            if not in_code:
                flush()
            buffer.append(line)
            in_code = not in_code
            if not in_code:
                flush()
            continue
        if in_code:
            buffer.append(line)
            continue

        is_table_line = stripped.startswith("|") and stripped.endswith("|")
        if is_table_line:
            if not in_table:
                flush()
                in_table = True
            buffer.append(line)
            continue
        if in_table:
            flush()
            in_table = False

        if not stripped:
            flush()
            continue
        buffer.append(line)

    flush()
    return units


def _iter_heading_sections(markdown: str) -> list[tuple[str, int, str]]:
    headings: list[tuple[str, int, int]] = []
    lines = markdown.splitlines()
    for index, line in enumerate(lines):
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            headings.append((match.group(2).strip(), len(match.group(1)), index))

    sections: list[tuple[str, int, str]] = []
    for index, (title, heading_level, start_line) in enumerate(headings):
        end_line = headings[index + 1][2] if index + 1 < len(headings) else len(lines)
        content = "\n".join(lines[start_line:end_line]).strip()
        if content:
            sections.append((title, heading_level, content))
    return sections


def _sections_metadata(markdown: str, docling_metadata: dict[str, Any]) -> list[dict[str, Any]]:
    toc = docling_metadata.get("table_of_contents") or []
    sections: list[dict[str, Any]] = []
    for item in toc:
        title = str(item.get("title") or "").strip()
        if not title:
            continue
        page_id = item.get("page_id")
        sections.append(
            {
                "title": title,
                "heading_level": int(item.get("heading_level") or 1),
                "page_start": int(page_id) + 1 if isinstance(page_id, int) else None,
            }
        )

    if sections:
        return sections

    return [
        {
            "title": title,
            "heading_level": heading_level,
            "page_start": None,
        }
        for title, heading_level, _ in _iter_heading_sections(markdown)
    ]


def _count_docling_blocks(data: Any) -> dict[str, int]:
    counter: Counter[str] = Counter()

    def visit(value: Any) -> None:
        if isinstance(value, dict):
            block_type = value.get("block_type") or value.get("label") or value.get("type")
            if block_type:
                counter[str(block_type)] += 1
            for child in value.values():
                visit(child)
        elif isinstance(value, list):
            for item in value:
                visit(item)

    visit(data)
    return dict(counter)


def _page_count(docling_json: Any, docling_metadata: dict[str, Any]) -> int:
    metadata_page_count = docling_metadata.get("page_count")
    if isinstance(metadata_page_count, int):
        return metadata_page_count

    page_stats = docling_metadata.get("page_stats")
    if isinstance(page_stats, list) and page_stats:
        return len(page_stats)

    if isinstance(docling_json, dict):
        pages = docling_json.get("pages")
        if isinstance(pages, dict | list):
            return len(pages)

    blocks = _count_docling_blocks(docling_json)
    if blocks.get("Page"):
        return int(blocks["Page"])
    if isinstance(docling_json, list):
        return len(docling_json)
    return 0


def _join_markdown_lines(lines: list[str]) -> str:
    output: list[str] = []
    paragraph: list[str] = []
    in_code = False

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            output.append(" ".join(paragraph).strip())
            paragraph = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            output.append(line.rstrip())
            in_code = not in_code
            continue
        if in_code:
            output.append(line.rstrip())
            continue
        if not stripped:
            flush_paragraph()
            if output and output[-1] != "":
                output.append("")
            continue
        if _is_markdown_structure(stripped):
            flush_paragraph()
            output.append(stripped)
            continue
        paragraph.append(stripped)

    flush_paragraph()
    return "\n".join(output)


def _is_markdown_structure(line: str) -> bool:
    return bool(
        line.startswith(("#", "|", ">", "-", "*", "+", "$$", "```"))
        or re.match(r"^\d+[.)]\s+", line)
    )


def _is_page_marker(line: str) -> bool:
    stripped = line.strip()
    return bool(
        re.match(r"^-?\s*\d+\s*-?$", stripped)
        or re.match(r"^page\s+\d+(\s+of\s+\d+)?$", stripped, re.I)
    )


def _noise_candidate(line: str) -> str:
    stripped = re.sub(r"\s+", " ", line.strip())
    if not stripped or len(stripped) < 3 or len(stripped) > 100:
        return ""
    if _is_markdown_structure(stripped) or _is_page_marker(stripped):
        return ""
    return stripped.casefold()


def _visible_text(text: str) -> str:
    return re.sub(r"[\W_]+", "", text, flags=re.UNICODE)


def _document_key(file_path: Path, file_hash: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "-", file_path.stem).strip("-").lower()
    return f"{stem}-{file_hash[:12]}"


def _normalize_title(title: str) -> str:
    return re.sub(r"\s+", " ", title.strip().casefold())


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, data: Any) -> None:
    path.write_text(
        json.dumps(_json_safe(data), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _model_to_data(value: Any) -> Any:
    if hasattr(value, "model_dump"):
        return value.model_dump()
    if hasattr(value, "dict"):
        return value.dict()
    return value


def _json_safe(value: Any) -> Any:
    value = _model_to_data(value)
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, list | tuple):
        return [_json_safe(item) for item in value]
    if isinstance(value, str | int | float | bool) or value is None:
        return value
    return str(value)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert local dataset documents into Docling Markdown/JSON/image artifacts.",
    )
    parser.add_argument(
        "--dataset-dir",
        type=Path,
        default=DEFAULT_DATASET_DIR,
        help="Folder containing source PDF/DOCX/PPTX files.",
    )
    parser.add_argument(
        "--processed-root",
        type=Path,
        default=DEFAULT_PROCESSED_DIR,
        help="Folder where processed artifacts will be written.",
    )
    parser.add_argument(
        "--dataset-name",
        default="docs",
        help="Dataset name stored in generated metadata.",
    )
    parser.add_argument(
        "--model-name",
        default="text-embedding-3-small",
        help="Tokenizer model name used for chunk preview metadata.",
    )
    parser.add_argument(
        "--convert-only",
        action="store_true",
        help="Only write raw.md, docling.json, docling_metadata.json, and images.",
    )
    parser.add_argument(
        "--no-cache",
        action="store_true",
        help="Re-run Docling even if raw artifacts already exist.",
    )
    args = parser.parse_args()

    document_files = collect_dataset_documents(args.dataset_dir)
    if not document_files:
        print({"status": "no_files", "dataset_dir": str(args.dataset_dir)})
        return 0

    seen_hashes: dict[str, Path] = {}
    failures = 0
    for document_file in document_files:
        try:
            validation = validate_document_file(document_file, seen_hashes=seen_hashes)
            processed_dir = args.processed_root / validation.document_key
            if args.convert_only:
                converted = convert_with_docling(
                    document_file,
                    processed_dir,
                    use_cache=not args.no_cache,
                )
                print(
                    {
                        "status": "converted",
                        "file": str(document_file),
                        "raw_md": str(converted.raw_markdown_path),
                        "raw_json": str(converted.raw_json_path),
                        "docling_metadata": str(converted.docling_metadata_path),
                        "images": converted.image_paths,
                    }
                )
                continue

            processed = process_document_for_rag(
                document_file,
                dataset_name=args.dataset_name,
                processed_root=args.processed_root,
                model_name=args.model_name,
                validation=validation,
                use_cache=not args.no_cache,
            )
            print(
                {
                    "status": "processed",
                    "file": str(document_file),
                    "processed_dir": str(processed.processed_dir),
                    "raw_md": str(processed.raw_markdown_path),
                    "raw_json": str(processed.raw_json_path),
                    "clean_md": str(processed.clean_markdown_path),
                    "metadata": str(processed.metadata_path),
                    "report": str(processed.report_path),
                    "chunks": len(processed.chunks),
                }
            )
        except DatasetPipelineError as exc:
            failures += 1
            print({"status": "failed", "file": str(document_file), "error": str(exc)})

    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
